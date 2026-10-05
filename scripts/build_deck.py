#!/usr/bin/env python3
"""Build the editable SpotStory pitch deck using the supplied Hackdays PDF layout."""
from pathlib import Path
import sys

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PITCH = ROOT / 'pitch'
ASSETS = PITCH / 'template_pages'
TEMPLATE = PITCH / 'template_source.pdf'
OUT = PITCH / 'SpotStory_Hackdays2O.pptx'

NAVY = '263956'
INK = '263632'
BODY = '506159'
MUTED = '87938E'
PALE = 'F1F3F8'
LINE = 'D2D8DE'
CORAL = 'E95B39'
BLUE = '318BE1'
AMBER = 'F3AE39'
GREEN = '2E6E58'
WHITE = 'FFFFFF'


def I(points):
    return Inches(points / 72)


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_backgrounds():
    ASSETS.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(TEMPLATE)
    for idx in range(7):
        out = ASSETS / f'template_{idx + 1}.png'
        pix = doc[idx].get_pixmap(matrix=fitz.Matrix(2.25, 2.25), alpha=False)
        pix.save(out)
    doc.close()


def add_bg(slide, page_num):
    slide.shapes.add_picture(str(ASSETS / f'template_{page_num}.png'), 0, 0, width=Inches(11), height=Inches(8.5))


def add_rect(slide, x, y, w, h, fill=PALE, line=LINE, radius=True, line_width=0.65):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_circle(slide, x, y, d, fill, text=None, text_color=WHITE, size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = I(0)
        tf.margin_top = tf.margin_bottom = I(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = 'Arial'
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(text_color)
    return shape


def add_text(slide, text, x, y, w, h, size=10, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font='Arial', margin=0,
             spacing=1.06, char_spacing=None):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(margin)
    tf.vertical_anchor = valign
    lines = str(text).split('\n')
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)
        if char_spacing is not None:
            r._r.get_or_add_rPr().set('spc', str(char_spacing))
    return box


def add_bullets(slide, bullets, x, y, w, h, size=9.3, color=BODY, gap=7, marker=CORAL):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(0)
    tf.margin_top = tf.margin_bottom = I(0)
    for idx, text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        p.level = 0
        marker_run = p.add_run()
        marker_run.text = '•  '
        marker_run.font.name = 'Arial'
        marker_run.font.size = Pt(size)
        marker_run.font.bold = True
        marker_run.font.color.rgb = rgb(marker)
        body_run = p.add_run()
        body_run.text = text
        body_run.font.name = 'Arial'
        body_run.font.size = Pt(size)
        body_run.font.color.rgb = rgb(color)
    return box


def cover_line(slide, y, text, color=MUTED):
    # Cover only the original template's instructional sentence; leave its page number and organizer mark intact.
    add_rect(slide, 21, y, 535, 17, fill=WHITE, line=None, radius=False)
    add_text(slide, text, 30, y + 2, 520, 14, size=7.1, color=color, italic=True)


def style_title_slide(slide):
    # Keep the event logos, brand marks, and labels from the supplied PDF; replace its editable placeholders.
    add_rect(slide, 31, 265, 490, 49, fill=WHITE, line=None, radius=False)
    add_text(slide, 'SpotStory', 36, 267, 445, 40, size=28, color=NAVY, bold=False)
    add_text(slide, 'Gemini-assisted skin-change visit prep · not cancer detection', 38, 309, 470, 17, size=9.0, color=GREEN)

    add_rect(slide, 31, 334, 426, 38, fill=WHITE, line=None, radius=False)
    add_text(slide, 'Team Name', 36, 335, 120, 12, size=7.0, color='6E7C8F')
    add_text(slide, '[TEAM NAME — exact Unstop name]', 36, 347, 386, 17, size=10.4, color=NAVY)

    add_rect(slide, 31, 408, 435, 27, fill=WHITE, line=None, radius=False)
    add_text(slide, '[Member 1]   ·   [Member 2]   ·   [Member 3]   ·   [Member 4]', 36, 412, 410, 17, size=8.2, color=NAVY)

    # Keep the official Tracks Chosen label but remove the unselected choices printed in the PDF.
    add_rect(slide, 27, 461, 548, 27, fill=WHITE, line=None, radius=False)
    chips = [('Healthcare', 37, CORAL, 70), ('AI / ML', 116, BLUE, 63), ('Open Innovation', 188, BLUE, 94)]
    for label, x, col, width in chips:
        add_rect(slide, x, 464, width, 19, fill=WHITE, line=col, radius=True, line_width=0.8)
        add_text(slide, label, x + 2, 468, width - 4, 12, size=6.9, color=col, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 29, 492, 516, 19, fill=WHITE, line=None, radius=False)
    add_text(slide, 'A conversation-prep prototype: visual notes, personal timeline, clinician handoff.', 36, 495, 498, 12, size=6.8, color=MUTED, italic=True)


def slide_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 2)
    # Cover the template card prompts, then rebuild the same two-column card layout.
    add_rect(slide, 27, 207, 344, 268, fill=PALE, line='C9CDD4', radius=True, line_width=0.75)
    add_rect(slide, 374, 207, 344, 268, fill=PALE, line='C9CDD4', radius=True, line_width=0.75)
    add_circle(slide, 49, 230, 27, CORAL, 'P', size=9)
    add_circle(slide, 395, 230, 27, BLUE, 'S', size=9)
    add_text(slide, 'The Problem', 86, 232, 238, 24, size=13, color=NAVY)
    add_text(slide, 'Our Solution', 432, 232, 238, 24, size=13, color=NAVY)
    add_bullets(slide, [
        'A single phone photo loses the timeline: when a spot appeared and how it changed.',
        'An unvalidated “AI cancer score” can overclaim what a camera image proves.',
        'ACS flags new/evolving spots; AAD includes all skin tones, palms, soles and nails.',
        'Need: an accessible bridge from “I noticed this” to a clearer care conversation.'
    ], 48, 269, 301, 184, size=9.3, gap=7)
    add_bullets(slide, [
        'SpotStory pairs a person’s own timeline with optional Gemini photo notes.',
        'Gemini describes visible cues only — no cancer label, probability, or “all clear.”',
        'The browser turns selected facts into an editable, visit-ready note.',
        'Static, sourced guidance supports clinician follow-up; no diagnosis or urgency score.'
    ], 395, 269, 300, 184, size=9.3, gap=7, marker=BLUE)
    cover_line(slide, 480, 'Sources: American Cancer Society · American Academy of Dermatology | Prototype, not clinically validated.', GREEN)
    return slide


def slide_stack(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 3)
    cards = [
        (27, 207, 245, 101, CORAL, 'Frontend', 'Vanilla HTML / CSS / JS\nResponsive, accessible UI\nLocal image preview + resize'),
        (280, 207, 242, 101, BLUE, 'Backend', 'Node.js 18 · built-in HTTP\nSame-origin Gemini proxy\nInput validation + rate limit'),
        (530, 207, 235, 101, AMBER, 'Data', 'No database\nNo image written to disk\nTimeline remains in browser'),
        (27, 315, 245, 101, CORAL, 'Public API', 'Google Gemini API\n`generateContent` image input\nServer-to-server request'),
        (280, 315, 242, 101, BLUE, 'AI / ML', 'Gemini 3.5 Flash\nStructured JSON visual notes\nNo diagnosis or risk score'),
        (530, 315, 235, 101, AMBER, 'Other Tools', 'Google AI Studio key\n`.env.local` server secret\nSynthetic/static demo mode')
    ]
    for x, y, w, h, accent, title, body in cards:
        add_rect(slide, x, y, w, h, fill=PALE, line='D0D4DA', radius=True, line_width=.65)
        add_circle(slide, x + 14, y + 13, 22, accent)
        add_text(slide, title, x + 45, y + 15, w - 55, 17, size=9.4, color=NAVY, bold=True)
        add_text(slide, body, x + 15, y + 45, w - 28, 48, size=8.3, color=BODY, spacing=1.14)
    cover_line(slide, 421, 'Live request sends only the resized image and selected output language; user timeline stays client-side.', GREEN)
    return slide


def slide_workflow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 4)
    # Replace the diagram placeholder inside the supplied panel.
    add_rect(slide, 27, 207, 738, 269, fill=PALE, line='C9CDD4', radius=True, line_width=.75)
    add_text(slide, 'LIVE FLOW · EXPLICIT CONSENT · NO PERSISTENT IMAGE STORE', 48, 228, 675, 15, size=7.5, color=MUTED, bold=True)
    boxes = [
        (49, BLUE, '01 · PERSON', 'Selects own\nspot history'),
        (187, CORAL, '02 · BROWSER', 'Previews +\nresizes photo'),
        (325, NAVY, '03 · NODE', 'Checks file;\nhides API key'),
        (463, GREEN, '04 · GEMINI', 'Returns JSON\nvisual notes'),
        (601, AMBER, '05 · HANDOFF', 'Local visit\nprep note')
    ]
    for x, accent, title, body in boxes:
        add_rect(slide, x, 279, 111, 79, fill=WHITE, line=accent, radius=True, line_width=1.0)
        add_text(slide, title, x + 7, 293, 97, 15, size=7.0, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 7, 315, 97, 31, size=8.6, color=INK, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, spacing=1.05)
    for x in [163, 301, 439, 577]:
        add_text(slide, '→', x, 303, 21, 25, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    # Replace old instructional sentence below the node row with clear data boundaries.
    add_text(slide, 'Only the photo + language reach Gemini', 65, 389, 200, 25, size=8.1, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 'Timeline + optional note stay in browser', 296, 389, 205, 25, size=8.1, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 'No diagnosis, probability, or urgency', 531, 389, 185, 25, size=8.1, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    cover_line(slide, 466, 'The photo is not written to disk by this prototype. Live Gemini requests are subject to Google API terms.', GREEN)
    return slide


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 5)
    add_rect(slide, 27, 207, 738, 269, fill=PALE, line='C9CDD4', radius=True, line_width=.75)
    add_text(slide, 'NO DATABASE · SAME-ORIGIN API · SECRET NEVER ENTERS THE BROWSER', 48, 228, 675, 15, size=7.5, color=MUTED, bold=True)
    layers = [
        (46, 160, BLUE, 'CLIENT', 'Browser UI', ['Vanilla JS', 'Story + consent', 'Image resized locally']),
        (226, 160, CORAL, 'APPLICATION', 'Node.js server', ['MIME / size check', 'Rate-limited proxy', 'Output safety filter']),
        (406, 160, GREEN, 'AI API', 'Gemini API', ['Image input only', 'Constrained JSON', 'Key in server header']),
        (586, 160, NAVY, 'DATA / INFRA', 'No DB', ['No file persistence', 'Secret in .env.local', 'Usage-based API'])
    ]
    for x, w, accent, label, title, details in layers:
        add_rect(slide, x, 267, w, 112, fill=WHITE, line=accent, radius=True, line_width=.9)
        add_text(slide, label, x + 12, 278, w - 24, 11, size=6.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 10, 295, w - 20, 20, size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, '\n'.join(details), x + 10, 324, w - 20, 48, size=7.3, color=BODY, align=PP_ALIGN.CENTER, spacing=1.16)
    for x in [207, 387, 567]:
        add_text(slide, '→', x, 308, 17, 24, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 'Returned notes + locally composed story → editable visit brief', 147, 407, 500, 19, size=8.4, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    cover_line(slide, 466, 'Future deployment would add identity, encryption, consent governance, clinical review, and privacy/regulatory assessment.', GREEN)
    return slide


def slide_usp(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 6)
    cards = [
        (26, CORAL, '1', 'Pixels → a visit brief', 'Gemini notes visible cues; the person supplies the timeline. The handoff — not a cancer score — is the product.'),
        (278, BLUE, '2', 'A different lane', 'Miiskin maps and compares photos; SkinVision offers risk indications. SpotStory explores a no-score, editable clinician brief.'),
        (530, AMBER, '3', 'A testable hypothesis', 'Bilingual, privacy-light, human-led. Clinician co-design and measured user value must prove it can stand out.')
    ]
    for x, accent, number, title, body in cards:
        add_rect(slide, x, 214, 239, 250, fill=PALE, line='D0D4DA', radius=True, line_width=.7)
        add_circle(slide, x + 101, 235, 34, accent, number, size=11)
        add_text(slide, title, x + 17, 286, 205, 31, size=13.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, body, x + 26, 326, 188, 83, size=9.2, color=BODY, align=PP_ALIGN.CENTER, spacing=1.17)
    cover_line(slide, 464, 'Landscape scan: Miiskin (photo tracking) · SkinVision (risk indication). Novelty is unproven; validate this no-score handoff with users.', GREEN)
    return slide


def slide_feasibility(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, 7)
    rows = [
        (211, CORAL, 'Technical Feasibility', 'Built: responsive UI, synthetic-image demo, server-side Gemini JSON endpoint, response filter, 10 safety/API stub tests, and local note builder.\nNot built: clinical validation, authentication, or production data controls.'),
        (294, BLUE, 'Market Feasibility', 'Product hypothesis: people tracking a changing spot may value a clearer visit note; clinicians may value concise patient context.\nNext: interview patients and dermatology/primary-care professionals; quantify willingness to adopt.'),
        (377, AMBER, 'Roadmap & Cost', '0–4 weeks: co-design + privacy/risk review. 1–3 months: usability + representative image testing. 3–6 months: supervised pilot only if approved.\nPrototype API spend: usage-based with a quota cap; production budget and regulatory path TBD.')
    ]
    for y, accent, title, body in rows:
        add_rect(slide, 27, y, 738, 70, fill=PALE, line='D0D4DA', radius=True, line_width=.65)
        add_circle(slide, 43, y + 20, 28, accent, '✓', size=11)
        add_text(slide, title, 83, y + 13, 210, 18, size=10, color=NAVY, bold=True)
        add_text(slide, body, 83, y + 34, 655, 31, size=7.8, color=BODY, spacing=1.12)
    cover_line(slide, 461, 'No patient-facing diagnostic use until expert review, representative validation, privacy/security work, and appropriate regulatory assessment.', GREEN)
    return slide


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f'Missing supplied template PDF: {TEMPLATE}')
    add_backgrounds()
    prs = Presentation()
    prs.slide_width = Inches(11)
    prs.slide_height = Inches(8.5)
    prs.core_properties.title = 'SpotStory — Hackdays 2.O'
    prs.core_properties.subject = 'Best Use of Google Gemini API | Non-diagnostic skin-change visit prep'
    prs.core_properties.author = 'SpotStory Hackdays Team'
    prs.core_properties.keywords = 'Hackdays 2.O, Gemini API, healthcare, AI, safety-by-design'

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(cover, 1)
    style_title_slide(cover)
    slide_problem(prs)
    slide_stack(prs)
    slide_workflow(prs)
    slide_architecture(prs)
    slide_usp(prs)
    slide_feasibility(prs)

    prs.save(OUT)
    print(f'Wrote {OUT}')
    print(f'Slides: {len(prs.slides)} | Size: 11 × 8.5 in | Supplied template style: yes')


if __name__ == '__main__':
    main()
